"""Synthetic integration contracts for the original dashboard and new pipeline."""

from __future__ import annotations

import json
from datetime import date
from io import BytesIO
from pathlib import Path

import pandas as pd
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from src.api.compat import parts as compat_parts
from src.api.compat.context import planning_context, published_order_hold_reason
from src.api.compat.exports import order_workbook
from src.api.compat.filters import clear_cache, mart
from src.api.main import app
from src.core.context import PlanningContext
from src.core.result import StageResult
from src.core.settings import get_settings
from src.dashboard import sales_abc, sku
from src.dashboard.compatibility import service_plan, validate_output
from src.io.parquet import read_table, write_table


@pytest.fixture
def workspace(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setenv("SPI_ROOT", str(tmp_path))
    get_settings.cache_clear()
    get_settings().ensure_dirs()
    clear_cache()
    yield tmp_path
    clear_cache()
    get_settings.cache_clear()


@pytest.fixture
def published_sku(workspace: Path) -> pd.DataFrame:
    key = {"active_sku_id": "SYNTHETIC-PART"}
    tables = {
        "sku_classification": {
            **key,
            "description": "Synthetic part",
            "abc": "A",
            "xyz": "X",
            "fsn": "F",
            "quadrant": "smooth",
            "mean_monthly_demand": 10.0,
            "cv_monthly": 0.2,
            "nonzero_months": 12,
            "months_observed": 12,
            "annual_consumption_value": 1200.0,
            "last_demand_month": "2025-12",
            "behaviour_class": "wear",
            "criticality": None,
        },
        "forecast_protection": {
            **key,
            "model": "naive",
            "mu_month": 10.0,
            "mu_month_baseline": 9.0,
            "mu_month_parc": 12.0,
            "mu_p": 40.0,
            "sigma_p": 4.0,
            "p50": 40.0,
            "p90": 45.0,
            "p95": 48.0,
        },
        "policy_params": {
            **key,
            "fill_target": 0.98,
            "z": 2.0,
            "safety_stock": 8.0,
            "ss_strategy": "normal",
            "rol": 38.0,
            "months_of_cover": 4.0,
            "unit_value": 12.0,
        },
        "policy_selection": {**key, "policy": "R_S"},
        "stock_position": {
            **key,
            "on_hand": 0.0,
            "on_order": 20.0,
            "backorders": 2.0,
            "ip": 18.0,
            "storage_locations": "SYNTHETIC",
        },
        "monthly_order_proposal": {
            **key,
            "q_final": 25.0,
            "q_raw": 22.0,
            "unit_cost": 10.0,
            "value": 250.0,
            "cycle_month": "2025-12",
            "lead_time_months": 4,
            "on_order_interpretation": "arrival",
            "incoming_orders_status": "VERIFIED",
            "stock_snapshot_as_of": pd.Timestamp("2025-11-30"),
            "on_orders_coverage_end": "2026-12",
            "on_orders_sha256": "a" * 64,
            "trigger_reason": "review",
        },
        "orders_by_material": {
            **key,
            "ordered_quantity": 120.0,
            "confirmed_quantity": 90.0,
            "order_value": 1200.0,
            "fill_rate": 0.75,
        },
        "part_master_enriched": {
            **key,
            "material": "SYNTHETIC-PART",
            "description": "Synthetic part",
            "chain_depth": 0,
            "compatible_models": "SYNTHETIC-MODEL",
            "part_kind": "shared",
            "material_type": "part",
            "material_group": "parts",
            "brand": "test",
        },
    }
    for name, row in tables.items():
        write_table(pd.DataFrame([row]), "facts", name)
    billed = pd.DataFrame(
        [
            {
                "material_code": "SYNTHETIC-PART",
                "material_description": "Synthetic part",
                "month": "2025-11",
                "Net Sales": 1000.0,
                "SlsVolQty": 1.0,
                "is_return": False,
            }
        ]
    )
    write_table(billed, "facts", "parts_sales")
    result = StageResult("15_dashboard")
    frame = sku.build(PlanningContext(as_of=date(2025, 12, 1)), result)
    validate_output(frame, "mart_ui_sku")
    write_table(frame, "marts", "mart_ui_sku")
    billed_abc, billed_audit = sales_abc.build(
        read_table("facts", "part_master_enriched"), billed, date(2025, 12, 1)
    )
    master_analysis = sku.build_part_master_analysis(frame, billed_abc)
    validate_output(master_analysis, "mart_ui_part_master_analysis")
    write_table(master_analysis, "marts", "mart_ui_part_master_analysis")
    write_table(billed_audit, "marts", "mart_ui_sales_abc_audit")
    write_table(sku.overview(frame, result), "marts", "mart_ui_overview")
    write_table(service_plan(frame), "marts", "mart_ui_service_plan")
    write_table(pd.DataFrame([tables["monthly_order_proposal"]]), "marts", "mart_monthly_order")
    return frame


def test_forecast_and_order_reconcile_to_new_pipeline(published_sku: pd.DataFrame) -> None:
    row = published_sku.iloc[0]
    assert row["forecast_lt"] == 40.0  # Four-month lead, not five-month protection demand.
    assert row["forecast_month"] == "2026-01"
    assert row["net_requirement"] == row["roq"] == 25.0  # Pack-rounded, not raw 22.
    assert row["unit_value_lkr"] == 10.0
    client = TestClient(app)
    overview = client.get("/api/v1/overview/kpis").json()
    assert overview["kpis"]["total_order_value_lkr"] == 250.0
    # Empty shelf but 20 units on order: awaiting stock, not a stockout.
    assert overview["stock_status"] == {"awaiting_stock": 1}
    policy = client.get("/api/v1/policy?urgency=soon").json()
    assert policy["rows"][0]["order_urgency"] == "soon"
    assert policy["rows"][0]["net_requirement"] == 25.0
    assert client.get("/api/v1/inventory?status=awaiting_stock").json()["total"] == 1
    assert client.get("/api/v1/inventory?status=stockout").json()["total"] == 0
    assert client.get("/api/v1/classification?tier=R_S&part_type=shared").json()["total"] == 1
    assert client.get("/api/v1/classification?part_type=absent").json()["total"] == 0


def test_analysis_uses_full_part_master_and_resolved_alias_search(
    published_sku: pd.DataFrame,
) -> None:
    master = read_table("facts", "part_master_enriched")
    extra = pd.DataFrame(
        [
            {
                "active_sku_id": "NEW-SYNTHETIC",
                "material": material,
                "description": description,
                "chain_depth": depth,
                "compatible_models": "MODEL-X",
                "part_kind": "shared",
                "material_group": "test-group",
                "material_type": "part",
                "brand": "YM",
            }
            for material, description, depth in (
                ("OLD-SYNTHETIC", "Old description", 1),
                ("NEW-SYNTHETIC", "Current description", 0),
            )
        ]
    )
    write_table(pd.concat([master, extra], ignore_index=True), "facts", "part_master_enriched")
    billed_abc, billed_audit = sales_abc.build(
        read_table("facts", "part_master_enriched"),
        read_table("facts", "parts_sales"),
        date(2025, 12, 1),
    )
    analysis = sku.build_part_master_analysis(published_sku, billed_abc)
    validate_output(analysis, "mart_ui_part_master_analysis")
    assert len(analysis) == 2
    write_table(analysis, "marts", "mart_ui_part_master_analysis")
    write_table(billed_audit, "marts", "mart_ui_sales_abc_audit")
    clear_cache()
    client = TestClient(app)
    classified = client.get("/api/v1/classification").json()
    assert classified["total"] == 2
    assert classified["classified_count"] == 1
    assert classified["unclassified_count"] == 1
    assert classified["active_count"] == 1
    assert classified["inactive_count"] == 1
    assert classified["active_count"] + classified["inactive_count"] == classified["total"]
    assert {r["sales_activity_12m"] for r in classified["rows"]} == {"ACTIVE", "INACTIVE"}
    assert classified["abc_fsn_counts"]["A"]["F"] == 1
    assert classified["classification_coverage"]["ABC"] == {"assigned": 1, "not_classified": 1}
    assert classified["classification_coverage"]["XYZ"] == {"assigned": 1, "not_classified": 1}
    assert all(
        counts["assigned"] + counts["not_classified"] == classified["total"]
        for counts in classified["classification_coverage"].values()
    )
    assert classified["sales_abc_audit"]["sales_lines"] == 1
    assert client.get("/api/v1/classification?scope=unclassified").json()["total"] == 1
    assert client.get("/api/v1/classification?scope=inactive").json()["total"] == 1
    assert client.get("/api/v1/classification?scope=active").json()["total"] == 1
    found = client.get("/api/v1/classification?search=OLD-SYNTHETIC").json()
    assert found["total"] == 1
    assert found["rows"][0]["active_sku_id"] == "NEW-SYNTHETIC"
    assert found["rows"][0]["description"] == "Current description"
    assert found["rows"][0]["has_planning"] is False
    assert found["rows"][0]["sales_activity_12m"] == "INACTIVE"
    assert found["rows"][0]["avg_monthly_demand"] is None
    assert found["rows"][0]["alias_count"] == 2
    inventory = client.get("/api/v1/inventory?status=not_assessed").json()
    assert inventory["total"] == 1
    assert inventory["rows"][0]["has_stock_snapshot"] is False
    assert inventory["rows"][0]["stock_on_hand"] is None
    assert inventory["rows"][0]["stock_value_lkr"] is None
    assert inventory["unassessed_count"] == 1


@pytest.mark.parametrize(
    "path",
    [
        "forecast",
        "policy",
        "classification",
        "inventory",
        "forecast/fused-demand",
        "inventory/position",
    ],
)
def test_pagination_validation(published_sku: pd.DataFrame, path: str) -> None:
    client = TestClient(app)
    assert client.get("/api/v1/" + path + "?offset=-1").status_code == 422
    assert client.get("/api/v1/" + path + "?limit=0").status_code == 422
    response = client.get("/api/v1/" + path + "?limit=1&offset=1")
    assert response.status_code == 200
    assert response.json()["total"] == 1
    assert response.json()["rows"] == []


def test_service_plan_uses_parc_and_keeps_final_policy(published_sku: pd.DataFrame) -> None:
    original = published_sku.copy(deep=True)
    frame = service_plan(published_sku)
    row = frame.loc[frame["horizon_months"] == 5].iloc[0]
    assert row["service_plan_qty"] == 60.0
    assert row["recommended_order"] == 25.0
    pd.testing.assert_frame_equal(published_sku, original)
    response = TestClient(app).get("/api/v1/policy/uio-service-plan?horizon_months=5").json()
    assert response["rows"][0]["service_plan_qty"] == 60.0
    assert response["rows"][0]["recommended_order"] == 25.0


def test_trend_does_not_mix_other_sku_returns(workspace: Path) -> None:
    write_table(
        pd.DataFrame(
            [
                {
                    "active_sku_id": "A",
                    "month": "2025-12",
                    "ordered_quantity": 10.0,
                    "confirmed_quantity": 8.0,
                    "order_value": 100.0,
                }
            ]
        ),
        "facts",
        "demand_history",
    )
    write_table(
        pd.DataFrame(
            [
                {"active_sku_id": "A", "month": "2025-12", "ordered_quantity": 2.0},
                {"active_sku_id": "B", "month": "2025-12", "ordered_quantity": 99.0},
            ]
        ),
        "facts",
        "returns_history",
    )
    write_table(
        pd.DataFrame(
            [
                {"active_sku_id": "A", "month": "2026-01", "forecast_quantity": 12.0},
                {"active_sku_id": "B", "month": "2026-01", "forecast_quantity": 99.0},
            ]
        ),
        "facts",
        "forecast_monthly_live",
    )
    rows = TestClient(app).get("/api/v1/forecast/trend?sku=A").json()
    actual = next(row for row in rows if not row["is_forecast"])
    future = next(row for row in rows if row["is_forecast"])
    assert actual["return_qty"] == 2.0
    assert actual["net_demand"] == 10.0  # C-order demand stays gross; returns are separate.
    assert future["year_month_str"] == "2026-01"
    assert future["forecast_qty"] == 12.0
    assert future["net_demand"] is None


def test_forecast_table_returns_each_part_through_next_year(
    published_sku: pd.DataFrame,
) -> None:
    months = [str(month) for month in pd.period_range("2026-10", "2027-12", freq="M")]
    write_table(
        pd.DataFrame(
            [
                {
                    "active_sku_id": sku,
                    "month": month,
                    "forecast_quantity": float(index + 1) if sku == "SYNTHETIC-PART" else 999.0,
                }
                for sku in ("SYNTHETIC-PART", "OTHER-PART")
                for index, month in enumerate(months)
            ]
        ),
        "facts",
        "forecast_monthly_live",
    )

    result = TestClient(app).get("/api/v1/forecast?limit=1").json()
    assert result["forecast_months"] == months
    assert result["forecast_month_count"] == 15
    assert result["rows"][0]["monthly_forecast"] == [float(index + 1) for index in range(15)]
    assert result["rows"][0]["monthly_forecast"][-1] == 15.0


def test_cache_observes_republished_mart(workspace: Path) -> None:
    write_table(pd.DataFrame({"value": [1]}), "marts", "synthetic")
    assert mart("synthetic").iloc[0]["value"] == 1
    write_table(pd.DataFrame({"value": [2, 3]}), "marts", "synthetic")
    assert mart("synthetic")["value"].tolist() == [2, 3]


def test_export_has_final_quantities_metadata_and_literal_cells(
    published_sku: pd.DataFrame,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr("src.api.compat.context._local_current_month", lambda: "2025-12")
    response = TestClient(app).get("/api/v1/policy/export.xlsx?urgency=soon")
    assert response.status_code == 200
    book = load_workbook(BytesIO(response.content))
    assert "Metadata" in book.sheetnames
    metadata = dict(book["Metadata"].values)
    assert {
        "source_file_hashes",
        "code_commit_sha",
        "generated_at_utc",
        "model_version",
    } <= metadata.keys()
    assert book["Order Plan"].oddFooter.left.text
    frame = pd.DataFrame(
        {"active_sku_id": ["001"], "description": ["=1+1"], "q_final": [25.0], "value": [250.0]}
    )
    book = load_workbook(BytesIO(order_workbook(frame, {"model_version": "synthetic"})))
    assert book["Order Plan"]["B2"].data_type == "s"


def test_old_published_order_cannot_be_exported(published_sku: pd.DataFrame) -> None:
    proposal = read_table("marts", "mart_monthly_order")
    proposal["lead_time_months"] = 3
    write_table(proposal, "marts", "mart_monthly_order")
    clear_cache()
    assert planning_context()["lead_time_months"] == 3
    assert "outdated import lead time" in (published_order_hold_reason() or "")
    response = TestClient(app).get("/api/v1/policy/export.xlsx")
    assert response.status_code == 409


def test_unverified_incoming_orders_cannot_be_exported(published_sku: pd.DataFrame) -> None:
    proposal = read_table("marts", "mart_monthly_order")
    proposal["incoming_orders_status"] = "UNVERIFIED"
    write_table(proposal, "marts", "mart_monthly_order")
    clear_cache()
    assert "Incoming orders" in (published_order_hold_reason() or "")
    assert TestClient(app).get("/api/v1/policy/export.xlsx").status_code == 409


def test_historical_order_cannot_be_exported(
    published_sku: pd.DataFrame, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr("src.api.compat.context._local_current_month", lambda: "2026-10")
    assert "past or future cycle" in (published_order_hold_reason() or "")
    assert TestClient(app).get("/api/v1/policy/export.xlsx").status_code == 409


def test_missing_api_does_not_serve_html(workspace: Path) -> None:
    client = TestClient(app)
    response = client.get("/api/v1/does-not-exist")
    assert response.status_code == 404
    assert response.headers["content-type"].startswith("application/json")
    assert client.get("/api/v1/overview/kpis").status_code == 503


def test_original_ui_is_served_and_deep_links_work(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    import src.api.main as main

    (tmp_path / "index.html").write_text("<html>Original UI</html>", encoding="utf-8")
    monkeypatch.setattr(main, "DASHBOARD_DIR", tmp_path)
    client = TestClient(app)
    assert client.get("/").text == "<html>Original UI</html>"
    assert client.get("/orders").text == "<html>Original UI</html>"


def test_run_history_reads_persisted_status(workspace: Path) -> None:
    report = {
        "started_at": "2025-12-01T00:00:00+00:00",
        "finished_at": "2025-12-01T00:01:00+00:00",
        "ok": False,
        "stages": [{"stage": "15_dashboard", "status": "failed"}],
    }
    (get_settings().reports_dir / "run_synthetic.json").write_text(
        json.dumps(report), encoding="utf-8"
    )
    client = TestClient(app)
    jobs = client.get("/api/v1/pipeline/jobs").json()
    assert jobs["total"] == 1
    assert jobs["jobs"][0]["status"] == "FAILED"
    assert client.get("/api/v1/pipeline/jobs/synthetic").json()["status"] == "FAILED"


def test_new_requirement_summary_covers_all_filtered_rows(published_sku: pd.DataFrame) -> None:
    extra = published_sku.copy()
    extra["material_9"] = extra["active_sku_id"] = "SYNTHETIC-SECOND"
    extra["description"] = "Second test part"
    extra["value"] = 400.0
    write_table(pd.concat([published_sku, extra], ignore_index=True), "marts", "mart_ui_sku")
    client = TestClient(app)
    data = client.get("/api/v1/policy?limit=1").json()
    assert len(data["rows"]) == 1
    assert data["total"] == 2
    assert data["total_order_value_lkr"] == 650.0
    filtered = client.get("/api/v1/policy?search=SECOND").json()
    assert filtered["total_order_value_lkr"] == 400.0
    assert filtered["total"] == 1
    forecast = client.get("/api/v1/forecast?search=SECOND").json()
    assert forecast["total"] == forecast["parc_skus"] == 1
    assert forecast["zero_demand_skus"] == 0


def test_default_service_plan_uses_protection_interval(published_sku: pd.DataFrame) -> None:
    write_table(
        pd.DataFrame([{"arm": "selected", "verdict": "FRONTIER", "fill_rate": 0.96}]),
        "marts",
        "mart_service_and_stock",
    )
    client = TestClient(app)
    data = client.get("/api/v1/policy/uio-service-plan").json()
    assert data["horizon_months"] == 5
    assert data["rows"][0]["service_plan_qty"] == 60.0
    assert data["rows"][0]["avg_monthly"] == 12.0
    overview = client.get("/api/v1/overview/kpis").json()
    assert overview["planning"]["policy_verdict"] == "FRONTIER"
    assert overview["planning"]["protection_interval_months"] == 5


def test_uio_comparison_does_not_copy_fleet_into_recent_sales(workspace: Path) -> None:
    write_table(
        pd.DataFrame(
            [
                {
                    "year": 2025,
                    "model_name": "Synthetic model",
                    "uio_forecast": 200.0,
                    "is_forecast": False,
                }
            ]
        ),
        "marts",
        "mart_ui_parc_by_model",
    )
    write_table(
        pd.DataFrame(
            [
                {
                    "year": 2025,
                    "model_name": "Synthetic model",
                    "units_sold": 12.0,
                    "returned": 1.0,
                    "revenue_lkr": 120.0,
                    "vins": 13.0,
                    "dealers": 1.0,
                }
            ]
        ),
        "marts",
        "mart_ui_mc_by_model",
    )
    data = TestClient(app).get("/api/v1/bikes/uio").json()
    assert data["external"][0]["uio"] == 200.0
    assert data["mcsi"][0]["uio"] == 12.0
    assert data["mcsi"][0]["uio_pct"] == 100.0


def test_master_view_uses_resolved_identity_without_catalogue(
    published_sku: pd.DataFrame, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(compat_parts, "_pn_yamaha_master_view", lambda *args: None)
    master = pd.DataFrame(
        [
            {
                "material": "OLD",
                "active_sku_id": "NEW",
                "description": "Old part",
                "brand": "YM",
                "chain_depth": 1,
                "compatible_models": None,
                "part_kind": None,
            },
            {
                "material": "NEW",
                "active_sku_id": "NEW",
                "description": "Current part",
                "brand": "YM",
                "chain_depth": 0,
                "compatible_models": None,
                "part_kind": None,
            },
        ]
    )
    write_table(master, "facts", "part_master_enriched")
    data = TestClient(app).get("/api/v1/parts/master-view").json()
    assert data["total"] == 1
    assert data["source"] == "PN_Yamaha"
    assert data["rows"][0]["part_no"] == "NEW"
    assert data["rows"][0]["description"] == "Current part"
    assert data["rows"][0]["kind"] == "unclassified"


def test_master_view_splits_current_brands_without_mc_models_on_obm(
    workspace: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(
        compat_parts, "_pn_yamaha_master_view", lambda *args: pytest.fail("raw YM view used")
    )
    monkeypatch.setattr(
        compat_parts,
        "_catalogue_compatibility",
        lambda master: ({}, {"source": "step_02", "error": "catalogue unavailable"}),
    )
    master = pd.DataFrame(
        [
            {
                "material": "OLD-MC",
                "active_sku_id": "CURRENT-OB",
                "brand": "YM",
                "chain_depth": 1,
                "description": "Old motorcycle number",
                "compatible_models": "MC MODEL",
            },
            {
                "material": "CURRENT-OB",
                "active_sku_id": "CURRENT-OB",
                "brand": "OB",
                "chain_depth": 0,
                "latest_ss": "CURRENT-OB",
                "1st Supersede": "OLD-MC",
                "description": "Outboard replacement",
                "compatible_models": "MC MODEL",
            },
            {
                "material": "CURRENT-MC",
                "active_sku_id": "CURRENT-MC",
                "brand": "YM",
                "chain_depth": 0,
                "latest_ss": "CURRENT-MC",
                "description": "Motorcycle part",
                "compatible_models": "MC MODEL",
            },
        ]
    )
    master["part_kind"] = "shared"
    master["material_group"] = "parts"
    write_table(master, "facts", "part_master_enriched")
    client = TestClient(app)
    mc = client.get("/api/v1/parts/master-view?segment=mc").json()
    obm = client.get("/api/v1/parts/master-view?segment=obm").json()
    assert (mc["total"], obm["total"]) == (1, 1)
    assert mc["rows"][0]["part_no"] == "CURRENT-MC"
    assert obm["rows"][0]["part_no"] == "CURRENT-OB"
    assert obm["rows"][0]["compatible_models"] == ""
    assert obm["total_models"] == 0
    assert obm["rows"][0]["latest_ss"] == "CURRENT-OB"
    assert obm["rows"][0]["supersedes"][0] == "OLD-MC"


def test_catalogue_models_follow_the_current_parts_product_type() -> None:
    master = pd.DataFrame(
        [
            {"material": "OLD-MC", "active_sku_id": "CURRENT-OB", "brand": "YM", "chain_depth": 1},
            {
                "material": "CURRENT-OB",
                "active_sku_id": "CURRENT-OB",
                "brand": "OB",
                "chain_depth": 0,
            },
            {
                "material": "CURRENT-MC",
                "active_sku_id": "CURRENT-MC",
                "brand": "YM",
                "chain_depth": 0,
            },
        ]
    )
    models, unmatched = compat_parts._models_by_active_sku(
        master,
        [
            ("OLD-MC", "MC", ["MC MODEL ON OLD NUMBER"]),
            ("CURRENT-OB", "OBM", ["OBM MODEL"]),
            ("CURRENT-MC", "MC", ["MC MODEL"]),
        ],
    )
    assert models == {"CURRENT-OB": {"OBM MODEL"}, "CURRENT-MC": {"MC MODEL"}}
    assert unmatched == 1


def test_location_unknown_values_are_not_reported_as_zero(workspace: Path) -> None:
    write_table(pd.DataFrame([{"Plant": "W1B4", "Unrestricted": 20.0}]), "facts", "stock_by_plant")
    row = TestClient(app).get("/api/v1/inventory/stock-by-location").json()[0]
    assert row["qty"] == 20.0
    assert row["value_lkr"] is None
    assert row["sku_count"] is None


def test_blank_descriptions_are_filled_from_the_chain_or_labelled(
    published_sku: pd.DataFrame,
) -> None:
    master = read_table("facts", "part_master_enriched")
    rows = [
        # Current number blank, older number described: take the older description.
        ("CHAIN-SYN", "CHAIN-SYN", None, 0),
        ("CHAIN-SYN", "CHAIN-OLD", "Described on the old number", 1),
        # No number in the chain has a description: label it, never drop it.
        ("BLANK-SYN", "BLANK-SYN", "  ", 0),
    ]
    extra = pd.DataFrame(
        [
            {
                "active_sku_id": sku_id,
                "material": material,
                "description": description,
                "chain_depth": depth,
                "compatible_models": "",
                "part_kind": "shared",
                "material_group": "test-group",
                "material_type": "part",
                "brand": "YM",
            }
            for sku_id, material, description, depth in rows
        ]
    )
    write_table(pd.concat([master, extra], ignore_index=True), "facts", "part_master_enriched")
    analysis = sku.build_part_master_analysis(published_sku)
    validate_output(analysis, "mart_ui_part_master_analysis")
    described = analysis.set_index("active_sku_id")["description"]
    assert described["CHAIN-SYN"] == "Described on the old number"
    assert described["BLANK-SYN"] == sku.MISSING_DESCRIPTION


def test_classification_download_matches_the_filtered_list(
    published_sku: pd.DataFrame,
) -> None:
    from io import BytesIO

    analysis = sku.build_part_master_analysis(published_sku)
    analysis["behaviour_class"] = ["wear part"] + ["engine part"] * (len(analysis) - 1)
    write_table(analysis, "marts", "mart_ui_part_master_analysis")
    clear_cache()
    client = TestClient(app)
    listed = client.get("/api/v1/classification?behaviour=wear%20part").json()
    assert listed["total"] == 1
    response = client.get("/api/v1/classification/export.xlsx?behaviour=wear%20part")
    assert response.status_code == 200
    parts = pd.read_excel(BytesIO(response.content), sheet_name="Parts", dtype=str)
    meta = pd.read_excel(BytesIO(response.content), sheet_name="Metadata", dtype=str)
    assert len(parts) == listed["total"]
    assert set(parts["Behaviour class"]) == {"wear part"}
    assert {"source_file_hashes", "code_commit_sha", "generated_at_utc", "filters"} <= set(
        meta["field"]
    )
