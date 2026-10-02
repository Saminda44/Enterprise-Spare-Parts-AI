"""Order plan audit, held-line recommendations and the buyer workbook; synthetic parts."""

from __future__ import annotations

import json
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook
from src.api.compat.exports import buyer_workbook
from src.dashboard.order_plan import audit_line, recent_demand, recommend_held


def _row(**values: object) -> pd.Series:
    base = {
        "recent_monthly_demand": 10.0,
        "q_final": 20.0,
        "forecast_month": 10.0,
        "stock_checked": 5.0,
        "active_months": 12,
        "q_review": 0.0,
        "fleet_forecast": None,
    }
    return pd.Series({**base, **values})


def test_recent_demand_rates_from_history() -> None:
    months = [f"2026-{m:02d}" for m in range(1, 13)]
    history = pd.DataFrame({"active_sku_id": "P1", "month": months, "ordered_quantity": 12.0})
    rates = recent_demand(history).set_index("active_sku_id").loc["P1"]
    assert rates["avg6"] == 12 and rates["avg12"] == 12
    assert rates["active_months"] == 12 and rates["last_order_month"] == "2026-12"


def test_a_line_in_line_with_recent_demand_needs_no_check() -> None:
    assert audit_line(_row()) == []


def test_audit_reasons() -> None:
    reasons = audit_line(
        _row(q_final=100.0, forecast_month=40.0, stock_checked=40.0, active_months=2)
    )
    text = "; ".join(reasons)
    assert "10.0 months of recent demand" in text
    assert "forecast 40/mo vs recent orders 10/mo" in text
    assert "already = 4.0 months" in text
    assert "only 2 month(s)" in text
    assert audit_line(_row(recent_monthly_demand=0.0))[0] == "no orders in the last 12 months"


def test_held_line_recommendations() -> None:
    # need = 10/mo x (1 + 3) - 5 on hand+order = 35
    assert recommend_held(_row(q_review=40.0), 3.0)[0] == "release"
    action, qty, why = recommend_held(_row(q_review=200.0), 3.0)
    assert (action, qty) == ("trim", 35.0) and "trim 200 to 35" in why
    assert recommend_held(_row(q_review=50.0, stock_checked=80.0), 3.0)[0] == "drop"
    assert recommend_held(_row(q_review=50.0, recent_monthly_demand=0.0), 3.0)[0] == "drop"
    fleet = recommend_held(_row(q_review=50.0, recent_monthly_demand=0.0, fleet_forecast=4.0), 3.0)
    assert fleet[0] == "review"


def test_buyer_workbook_sheets_and_readable_headings() -> None:
    incoming = json.dumps([{"month": "2026-10", "qty": 50.0}])
    plan = pd.DataFrame(
        [
            {
                "active_sku_id": "P1",
                "description": "BOLT",
                "status": "to order",
                "abc": "A",
                "q_final": 20.0,
                "unit_cost": 10.0,
                "value": 200.0,
                "expected_arrival": "2027-01",
                "forecast_month": 10.0,
                "safety_stock": 5.0,
                "reorder_level": 15.0,
                "stock_checked": 5.0,
                "needs_check": False,
                "check_reasons": "",
                "q_review": 0.0,
                "value_review": 0.0,
                "recommendation": None,
                "suggested_qty": None,
                "suggested_value": None,
                "recommendation_reason": None,
                "recent_monthly_demand": 10.0,
                "expedite": True,
                "on_hand": 0.0,
                "incoming_by_month": incoming,
                "run_out_month": "2026-09",
            },
            {
                "active_sku_id": "P2",
                "description": "NUT",
                "status": "held for review",
                "abc": "C",
                "q_final": 0.0,
                "unit_cost": 5.0,
                "value": 0.0,
                "expected_arrival": "2027-01",
                "forecast_month": 1.0,
                "safety_stock": 3.0,
                "reorder_level": 4.0,
                "stock_checked": 0.0,
                "needs_check": False,
                "check_reasons": "",
                "q_review": 90.0,
                "value_review": 450.0,
                "recommendation": "trim",
                "suggested_qty": 4.0,
                "suggested_value": 20.0,
                "recommendation_reason": "trim 90 to 4",
                "recent_monthly_demand": 1.0,
                "expedite": False,
                "on_hand": 0.0,
                "incoming_by_month": "[]",
                "run_out_month": None,
            },
        ]
    )
    book = load_workbook(BytesIO(buyer_workbook(plan, {"model_version": "synthetic"})))
    assert book.sheetnames == ["Order", "Buyer review", "Expedite", "Metadata"]
    order = [c.value for c in book["Order"][1]]
    assert order[:6] == [
        "Part no.",
        "Description",
        "ABC",
        "Order qty",
        "Unit cost (LKR)",
        "Value (LKR)",
    ]
    assert book["Order"]["A2"].value == "P1" and book["Buyer review"]["F2"].value == "trim"
    assert book["Expedite"]["F2"].value == "2026-10: 50"
