"""Read each workbook once, convert to parquet, and never open it again.

orders.xlsx is 47 MB and MCSI.xlsx 11 MB; re-reading per stage costs minutes per run.
Conversion is keyed on the file's SHA-256, so a re-run is a no-op and a re-dropped file
with changed contents is picked up automatically.
"""

from __future__ import annotations

import hashlib
import json
import re
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import pandas as pd
import pyarrow as pa

from src.core.errors import SourceDataError
from src.core.settings import get_settings

_CHUNK = 1 << 20


def _as_text(value: Any) -> str | None:
    """Stringify a scalar, keeping nulls null and not turning 26198 into '26198.0'."""
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def coerce_for_parquet(frame: pd.DataFrame) -> tuple[pd.DataFrame, list[str]]:
    """Cast object columns that pyarrow refuses — SAP exports mix ints and strings.

    Business meaning: ``Dealer Code`` arrives as 26198 on some rows and "26198" on
    others. Stored as-is the column cannot be written at all, and joined as-is it
    silently misses. Raw mirrors keep the value as text; later steps parse deliberately.
    Coerced columns are recorded in the sidecar so this is never invisible.
    """
    coerced: list[str] = []
    out = frame
    for column in frame.columns:
        if frame[column].dtype != object:
            continue
        try:
            pa.array(frame[column], from_pandas=True)
        except (pa.ArrowInvalid, pa.ArrowTypeError, pa.ArrowNotImplementedError):
            if out is frame:
                out = frame.copy()
            out[column] = frame[column].map(_as_text)
            coerced.append(str(column))
    return out, coerced


def file_digest(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as fh:
        while chunk := fh.read(_CHUNK):
            digest.update(chunk)
    return digest.hexdigest()


def _slug(value: str) -> str:
    return re.sub(r"[^0-9a-z]+", "_", value.lower()).strip("_")


@dataclass
class ConversionResult:
    """What one workbook→parquet conversion produced."""

    source: Path
    parquet: Path
    sha256: str
    rows: int
    columns: int
    sheet: str
    converted: bool = True
    coerced_to_text: list[str] = field(default_factory=list)
    meta: dict[str, Any] = field(default_factory=dict)


def _meta_path(parquet: Path) -> Path:
    return parquet.with_suffix(".meta.json")


def workbook_to_parquet(
    source: Path,
    *,
    sheet: str | int = 0,
    out_dir: Path | None = None,
    force: bool = False,
) -> ConversionResult:
    """Convert one sheet of one workbook to parquet, skipping unchanged files.

    The sidecar ``*.meta.json`` records the source hash, row and column counts and the
    ingestion time — the vintage of the data every downstream number rests on.
    """
    source = Path(source)
    if not source.exists():
        raise SourceDataError(f"source workbook not found: {source}")

    settings = get_settings()
    out_dir = out_dir or settings.source_parquet_dir
    out_dir.mkdir(parents=True, exist_ok=True)

    stem = _slug(source.stem)
    name = stem if sheet in (0, None) else f"{stem}__{_slug(str(sheet))}"
    parquet = out_dir / f"{name}.parquet"
    meta_path = _meta_path(parquet)

    digest = file_digest(source)
    if parquet.exists() and meta_path.exists() and not force:
        previous: dict[str, Any] = json.loads(meta_path.read_text(encoding="utf-8"))
        if previous.get("sha256") == digest:
            return ConversionResult(
                source=source,
                parquet=parquet,
                sha256=digest,
                rows=int(previous.get("rows", 0)),
                columns=int(previous.get("columns", 0)),
                sheet=str(previous.get("sheet", sheet)),
                converted=False,
                coerced_to_text=list(previous.get("coerced_to_text", [])),
                meta=previous,
            )

    frame = pd.read_excel(source, sheet_name=sheet)
    frame, coerced = coerce_for_parquet(frame)
    frame.to_parquet(parquet, index=False)

    rows = int(len(frame))
    columns = int(frame.shape[1])
    meta: dict[str, Any] = {
        "source": source.name,
        "source_path": str(source),
        "sheet": str(sheet),
        "sha256": digest,
        "rows": rows,
        "columns": columns,
        "coerced_to_text": coerced,
        "source_bytes": source.stat().st_size,
        "source_modified": datetime.fromtimestamp(source.stat().st_mtime, tz=UTC).isoformat(
            timespec="seconds"
        ),
        "ingested_at": datetime.now(UTC).isoformat(timespec="seconds"),
    }
    meta_path.write_text(json.dumps(meta, indent=2), encoding="utf-8")

    return ConversionResult(
        source=source,
        parquet=parquet,
        sha256=digest,
        rows=rows,
        columns=columns,
        sheet=str(sheet),
        converted=True,
        coerced_to_text=coerced,
        meta=meta,
    )


def workbook_sheets(source: Path) -> list[str]:
    """Sheet names without loading any data."""
    from openpyxl import load_workbook

    book = load_workbook(source, read_only=True)
    try:
        return list(book.sheetnames)
    finally:
        book.close()


def read_source(name: str, *, out_dir: Path | None = None) -> pd.DataFrame:
    """Read a previously converted source by its slug (e.g. ``"orders"``)."""
    return pd.read_parquet(_resolve_source(name, out_dir))


def _resolve_source(name: str, out_dir: Path | None) -> Path:
    """Accept either the exact stem or a sluggable name.

    ``sales_summery__model_classification`` is already a stem; re-slugging it would
    collapse the ``__`` sheet separator and miss the file.
    """
    settings = get_settings()
    out_dir = out_dir or settings.source_parquet_dir
    for candidate in (out_dir / f"{name}.parquet", out_dir / f"{_slug(name)}.parquet"):
        if candidate.exists():
            return candidate
    raise SourceDataError(f"{name}.parquet not found — convert the workbook first")


def source_vintage(name: str, *, out_dir: Path | None = None) -> dict[str, Any]:
    """The recorded hash, row count and ingestion time for a converted source."""
    meta_path = _meta_path(_resolve_source(name, out_dir))
    if not meta_path.exists():
        raise SourceDataError(f"no vintage recorded for {name!r}")
    vintage: dict[str, Any] = json.loads(meta_path.read_text(encoding="utf-8"))
    return vintage
